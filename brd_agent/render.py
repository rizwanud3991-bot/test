"""Turn structured outputs into Markdown, CSV, HTML slide decks, PowerPoint and the portal page."""

from __future__ import annotations

import csv
import html
import json
from pathlib import Path

from .schemas import Deck, QualityReview, Requirement, RequirementsAnalysis, TraceabilityMatrix


def _bullets(items: list[str]) -> str:
    return "\n".join(f"- {item}" for item in items) if items else "_None identified._"


def _cell(text: str) -> str:
    return str(text).replace("|", "\\|").replace("\n", " ")


def _req_table(reqs: list[Requirement]) -> str:
    if not reqs:
        return "_None identified._"
    lines = [
        "| ID | Title | Priority | Category | Source |",
        "|---|---|---|---|---|",
    ]
    lines += [
        f"| {r.id} | {_cell(r.title)} | {r.priority} | {_cell(r.category)} | {_cell(r.source)} |"
        for r in reqs
    ]
    details = []
    for r in reqs:
        criteria = "\n".join(f"  - {c}" for c in r.acceptance_criteria)
        details.append(f"**{r.id} — {r.title}**\n\n{r.description}\n\n- Acceptance criteria:\n{criteria}")
    return "\n".join(lines) + "\n\n" + "\n\n".join(details)


def analysis_markdown(a: RequirementsAnalysis) -> str:
    stakeholders = "\n".join(
        f"| {_cell(s.name)} | {_cell(s.interest)} | {s.influence} |" for s in a.stakeholders
    )
    return f"""# Requirements Analysis — {a.project_name}

## Problem statement
{a.problem_statement}

## Business objectives
{_bullets(a.business_objectives)}

## Success metrics (KPIs)
{_bullets(a.success_metrics)}

## Stakeholders
| Stakeholder | Interest | Influence |
|---|---|---|
{stakeholders}

## Scope
**In scope**
{_bullets(a.in_scope)}

**Out of scope**
{_bullets(a.out_of_scope)}

## Functional requirements
{_req_table(a.functional_requirements)}

## Non-functional requirements
{_req_table(a.non_functional_requirements)}

## Constraints
{_bullets(a.constraints)}

## Assumptions
{_bullets(a.assumptions)}

## Dependencies
{_bullets(a.dependencies)}
"""


def gaps_markdown(a: RequirementsAnalysis) -> str:
    gaps = "\n".join(
        f"| {i} | {_cell(g.area)} | {_cell(g.issue)} | {g.impact} | {_cell(g.clarifying_question)} |"
        for i, g in enumerate(a.gaps_and_ambiguities, start=1)
    )
    ideas = "\n".join(
        f"| {_cell(s.title)} | {_cell(s.rationale)} | {_cell(s.benefit_for)} | {s.effort} |"
        for s in a.suggested_enhancements
    )
    risks = "\n".join(
        f"| {r.id} | {_cell(r.description)} | {r.likelihood} | {r.impact} | {_cell(r.mitigation)} |"
        for r in a.risks
    )
    glossary = "\n".join(f"| {_cell(t.term)} | {_cell(t.meaning)} |" for t in a.glossary)
    return f"""# Gaps, Clarifying Questions & Ideas — {a.project_name}

Send the questions below to the business owner before design sign-off. Each one closes a gap that
could otherwise change scope, cost or design.

## Gaps and clarifying questions
| # | Area | Issue | Impact | Question for the business |
|---|---|---|---|---|
{gaps}

## Suggested enhancements (not in the BRD, but worth considering)
| Idea | Why it adds value | Who benefits | Effort |
|---|---|---|---|
{ideas}

## Initial risk register
| ID | Risk | Likelihood | Impact | Mitigation |
|---|---|---|---|---|
{risks}

## Glossary
| Term | Meaning |
|---|---|
{glossary}
"""


def traceability_markdown(m: TraceabilityMatrix) -> str:
    rows = "\n".join(
        f"| {r.requirement_id} | {_cell(r.requirement_title)} | {_cell(r.hld_component)} | "
        f"{_cell(r.lld_element)} | {_cell(r.test_reference)} | {_cell(r.delivery_phase)} | "
        f"{r.status} | {_cell(r.note)} |"
        for r in m.rows
    )
    covered = sum(r.status == "Covered" for r in m.rows)
    total = len(m.rows) or 1
    return f"""# Requirements Traceability Matrix

Coverage: **{covered}/{len(m.rows)} requirements fully covered ({covered * 100 // total}%)**

| Req ID | Requirement | HLD component | LLD element | Tests | Phase | Status | Note |
|---|---|---|---|---|---|---|---|
{rows}
"""


def write_traceability_csv(m: TraceabilityMatrix, path: Path) -> None:
    fields = list(m.rows[0].model_dump().keys()) if m.rows else ["requirement_id"]
    with path.open("w", newline="", encoding="utf-8") as f:
        writer = csv.DictWriter(f, fieldnames=fields)
        writer.writeheader()
        for row in m.rows:
            writer.writerow(row.model_dump())


def review_markdown(r: QualityReview) -> str:
    scores = "\n".join(f"| {_cell(s.dimension)} | {s.score}/5 | {_cell(s.comment)} |" for s in r.scores)
    return f"""# Quality Review of the Solution Pack

**Overall score: {r.overall_score}/5** — {r.verdict}

| Dimension | Score | Comment |
|---|---|---|
{scores}

## Requirements not adequately covered
{_bullets(r.uncovered_requirements)}

## Inconsistencies between documents
{_bullets(r.inconsistencies)}

## Recommended fixes (most important first)
{_bullets(r.recommended_fixes)}
"""


def deck_markdown(d: Deck) -> str:
    parts = [f"# {d.title}\n\n_Audience: {d.audience}_"]
    for i, s in enumerate(d.slides, start=1):
        sub = f"\n_{s.subtitle}_\n" if s.subtitle else ""
        parts.append(f"## {i}. {s.title}\n{sub}\n{_bullets(s.bullets)}\n\n> Speaker notes: {s.speaker_notes}")
    return "\n\n---\n\n".join(parts) + "\n"


DECK_TEMPLATE = """<!doctype html>
<html lang="en">
<head>
<meta charset="utf-8">
<meta name="viewport" content="width=device-width, initial-scale=1">
<title>__TITLE__</title>
<style>
  :root { --bg:#0f172a; --card:#ffffff; --ink:#0f172a; --muted:#475569; --accent:#2563eb; }
  * { box-sizing: border-box; }
  body { margin:0; background:var(--bg); font-family: system-ui, -apple-system, "Segoe UI", sans-serif; color:var(--ink); }
  .slide { display:none; width:min(1100px, 100vw - 32px); aspect-ratio:16/9; margin:24px auto; background:var(--card);
           border-radius:14px; padding:clamp(20px,5vw,64px); position:relative; overflow:auto; }
  .slide.active { display:block; }
  .slide h1 { font-size:clamp(22px,3.4vw,40px); margin:0 0 8px; }
  .slide h2 { font-size:clamp(14px,1.8vw,20px); color:var(--accent); font-weight:500; margin:0 0 28px; }
  .slide ul { font-size:clamp(14px,1.9vw,22px); line-height:1.55; padding-left:1.2em; }
  .slide.title-slide { display:none; background:linear-gradient(135deg,#1e3a8a,#2563eb); color:#fff; }
  .slide.title-slide.active { display:flex; flex-direction:column; justify-content:center; }
  .slide.title-slide h2 { color:#dbeafe; }
  .num { position:absolute; bottom:16px; right:24px; color:var(--muted); font-size:13px; }
  .notes { display:none; width:min(1100px, 100vw - 32px); margin:0 auto 24px; background:#1e293b; color:#e2e8f0;
           padding:16px 20px; border-radius:10px; line-height:1.5; }
  body.show-notes .notes.active { display:block; }
  .bar { text-align:center; color:#94a3b8; font-size:13px; padding-bottom:16px; }
  button { background:#1e293b; color:#e2e8f0; border:1px solid #334155; border-radius:8px; padding:6px 14px; margin:0 4px; cursor:pointer; }
  @media print {
    body { background:#fff; } .bar, .notes { display:none !important; }
    .slide, .slide.title-slide { display:block !important; page-break-after:always; margin:0; border-radius:0; width:100%; }
  }
</style>
</head>
<body>
__SLIDES__
<div class="bar">
  <button onclick="go(-1)">&larr; Prev</button>
  <button onclick="document.body.classList.toggle('show-notes')">Speaker notes (N)</button>
  <button onclick="go(1)">Next &rarr;</button>
  <div style="margin-top:8px">Arrow keys to navigate · N for notes · Ctrl/Cmd+P to save as PDF</div>
</div>
<script>
  const slides = [...document.querySelectorAll('.slide')];
  const notes = [...document.querySelectorAll('.notes')];
  let i = 0;
  function show() {
    slides.forEach((s, k) => s.classList.toggle('active', k === i));
    notes.forEach((n, k) => n.classList.toggle('active', k === i));
  }
  function go(d) { i = Math.max(0, Math.min(slides.length - 1, i + d)); show(); }
  document.addEventListener('keydown', e => {
    if (['ArrowRight', 'PageDown', ' '].includes(e.key)) go(1);
    if (['ArrowLeft', 'PageUp'].includes(e.key)) go(-1);
    if (e.key.toLowerCase() === 'n') document.body.classList.toggle('show-notes');
  });
  show();
</script>
</body>
</html>
"""


def deck_html(d: Deck) -> str:
    e = html.escape
    parts = []
    for i, s in enumerate(d.slides, start=1):
        cls = "slide title-slide" if i == 1 else "slide"
        bullets = "".join(f"<li>{e(b)}</li>" for b in s.bullets)
        sub = f"<h2>{e(s.subtitle)}</h2>" if s.subtitle else ""
        parts.append(
            f'<section class="{cls}"><h1>{e(s.title)}</h1>{sub}<ul>{bullets}</ul>'
            f'<div class="num">{i} / {len(d.slides)}</div></section>'
            f'<div class="notes"><strong>Speaker notes:</strong> {e(s.speaker_notes)}</div>'
        )
    return DECK_TEMPLATE.replace("__TITLE__", e(d.title)).replace("__SLIDES__", "\n".join(parts))


def write_pptx(d: Deck, path: Path) -> bool:
    """Write a PowerPoint version of the deck. Returns False if python-pptx is not installed."""
    try:
        from pptx import Presentation
        from pptx.util import Pt
    except ImportError:
        return False
    prs = Presentation()
    for i, s in enumerate(d.slides):
        if i == 0:
            slide = prs.slides.add_slide(prs.slide_layouts[0])
            slide.shapes.title.text = s.title
            slide.placeholders[1].text = s.subtitle or d.audience
        else:
            slide = prs.slides.add_slide(prs.slide_layouts[1])
            slide.shapes.title.text = s.title
            body = slide.placeholders[1].text_frame
            body.clear()
            lines = ([s.subtitle] if s.subtitle else []) + s.bullets
            for k, line in enumerate(lines):
                para = body.paragraphs[0] if k == 0 else body.add_paragraph()
                para.text = line
                para.font.size = Pt(20)
                if s.subtitle and k == 0:
                    para.font.bold = True
        slide.notes_slide.notes_text_frame.text = s.speaker_notes
    prs.save(str(path))
    return True


PORTAL_TEMPLATE = """<!doctype html>
<html lang="en">
<head>
<meta charset="utf-8">
<meta name="viewport" content="width=device-width, initial-scale=1">
<title>Solution Pack</title>
<style>
  :root { --bg:#f8fafc; --panel:#fff; --ink:#0f172a; --muted:#64748b; --line:#e2e8f0; --accent:#2563eb; }
  @media (prefers-color-scheme: dark) {
    :root { --bg:#0b1120; --panel:#111827; --ink:#e5e7eb; --muted:#94a3b8; --line:#1f2937; --accent:#60a5fa; }
  }
  * { box-sizing:border-box; }
  body { margin:0; background:var(--bg); color:var(--ink); font-family:system-ui,-apple-system,"Segoe UI",sans-serif; }
  header { padding:20px 16px; border-bottom:1px solid var(--line); }
  header h1 { margin:0; font-size:22px; }
  header p { margin:4px 0 0; color:var(--muted); font-size:14px; }
  .views { display:flex; gap:8px; flex-wrap:wrap; margin-top:12px; }
  .views button { border:1px solid var(--line); background:var(--panel); color:var(--ink); padding:6px 12px;
                  border-radius:999px; cursor:pointer; font-size:14px; }
  .views button.on { background:var(--accent); border-color:var(--accent); color:#fff; }
  .wrap { display:grid; grid-template-columns:260px 1fr; min-height:calc(100vh - 120px); }
  nav { border-right:1px solid var(--line); padding:12px; }
  nav a { display:block; padding:8px 10px; border-radius:8px; color:var(--ink); text-decoration:none; font-size:14px; }
  nav a.on { background:var(--accent); color:#fff; }
  nav a small { display:block; color:inherit; opacity:.7; }
  main { padding:16px clamp(16px,4vw,48px); max-width:1100px; overflow-x:auto; }
  main table { border-collapse:collapse; display:block; overflow-x:auto; font-size:14px; }
  main th, main td { border:1px solid var(--line); padding:6px 8px; vertical-align:top; }
  main pre { background:var(--panel); padding:12px; overflow-x:auto; border-radius:8px; }
  main blockquote { border-left:4px solid var(--accent); margin:0; padding:4px 16px; background:var(--panel); }
  .links a { margin-right:12px; color:var(--accent); }
  @media (max-width: 760px) { .wrap { grid-template-columns:1fr; } nav { border-right:0; border-bottom:1px solid var(--line); } }
</style>
</head>
<body>
<header>
  <h1 id="project"></h1>
  <p id="meta"></p>
  <div class="views" id="views"></div>
</header>
<div class="wrap">
  <nav id="nav"></nav>
  <main id="doc"></main>
</div>
<script id="pack" type="application/json">__PACK__</script>
<script src="https://cdn.jsdelivr.net/npm/marked@12.0.2/marked.min.js"></script>
<script src="https://cdn.jsdelivr.net/npm/mermaid@10.9.1/dist/mermaid.min.js"></script>
<script>
  const pack = JSON.parse(document.getElementById('pack').textContent);
  const views = {
    executive: { label: 'Executive view', docs: pack.views.executive },
    business: { label: 'Business teams view', docs: pack.views.business },
    technical: { label: 'Technical view', docs: pack.views.technical },
    all: { label: 'Everything', docs: pack.docs.map(d => d.file) },
  };
  let view = 'executive', current = null;
  document.getElementById('project').textContent = pack.project;
  document.getElementById('meta').textContent = pack.meta;
  if (window.mermaid) mermaid.initialize({ startOnLoad: false, securityLevel: 'strict',
    theme: matchMedia('(prefers-color-scheme: dark)').matches ? 'dark' : 'default' });

  function renderViews() {
    const box = document.getElementById('views');
    box.innerHTML = '';
    for (const [key, v] of Object.entries(views)) {
      const b = document.createElement('button');
      b.textContent = v.label; b.className = key === view ? 'on' : '';
      b.onclick = () => { view = key; current = null; renderViews(); renderNav(); };
      box.appendChild(b);
    }
  }
  function renderNav() {
    const nav = document.getElementById('nav');
    nav.innerHTML = '';
    const docs = pack.docs.filter(d => views[view].docs.includes(d.file));
    if (!current || !docs.some(d => d.file === current)) current = docs.length ? docs[0].file : null;
    for (const d of docs) {
      const a = document.createElement('a');
      a.href = '#'; a.className = d.file === current ? 'on' : '';
      a.innerHTML = '';
      a.append(d.title);
      const s = document.createElement('small'); s.textContent = d.audience; a.append(s);
      a.onclick = e => { e.preventDefault(); current = d.file; renderNav(); };
      nav.appendChild(a);
    }
    renderDoc();
  }
  async function renderDoc() {
    const main = document.getElementById('doc');
    const d = pack.docs.find(x => x.file === current);
    if (!d) { main.textContent = 'No documents yet.'; return; }
    const extra = (d.links || []).map(l => `<a href="${l}" target="_blank" rel="noopener">${l}</a>`).join('');
    if (window.marked) {
      main.innerHTML = (extra ? `<p class="links">${extra}</p>` : '') + marked.parse(d.markdown);
      main.querySelectorAll('pre code.language-mermaid').forEach(code => {
        const div = document.createElement('div');
        div.className = 'mermaid'; div.textContent = code.textContent;
        code.parentElement.replaceWith(div);
      });
      if (window.mermaid) { try { await mermaid.run({ nodes: main.querySelectorAll('.mermaid') }); } catch (e) {} }
    } else {
      const pre = document.createElement('pre'); pre.textContent = d.markdown;
      main.replaceChildren(pre);
    }
    main.scrollTop = 0;
  }
  renderViews(); renderNav();
</script>
</body>
</html>
"""


def portal_html(project: str, meta: str, docs: list[dict], views: dict[str, list[str]]) -> str:
    payload = json.dumps({"project": project, "meta": meta, "docs": docs, "views": views})
    # Keep embedded Markdown from closing the <script> tag early.
    payload = payload.replace("</", "<\\/")
    return PORTAL_TEMPLATE.replace("__PACK__", payload)
