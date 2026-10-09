"""Load business requirement documents from a folder."""

from __future__ import annotations

import hashlib
from dataclasses import dataclass
from pathlib import Path

SUPPORTED_EXTENSIONS = {".md", ".markdown", ".txt", ".rst", ".csv", ".json", ".yaml", ".yml", ".pdf", ".docx"}


@dataclass
class SourceDocument:
    path: Path
    name: str
    text: str

    @property
    def sha256(self) -> str:
        return hashlib.sha256(self.text.encode("utf-8")).hexdigest()


def _read_pdf(path: Path) -> str:
    try:
        from pypdf import PdfReader
    except ImportError as e:  # pragma: no cover - depends on environment
        raise RuntimeError(f"Reading {path.name} needs pypdf: pip install pypdf") from e
    reader = PdfReader(str(path))
    pages = []
    for i, page in enumerate(reader.pages, start=1):
        pages.append(f"[Page {i}]\n{page.extract_text() or ''}")
    return "\n\n".join(pages)


def _read_docx(path: Path) -> str:
    try:
        import docx
    except ImportError as e:  # pragma: no cover - depends on environment
        raise RuntimeError(f"Reading {path.name} needs python-docx: pip install python-docx") from e
    document = docx.Document(str(path))
    parts: list[str] = []
    for para in document.paragraphs:
        text = para.text.strip()
        if not text:
            continue
        style = (para.style.name or "").lower() if para.style is not None else ""
        if style.startswith("heading"):
            level = "".join(ch for ch in style if ch.isdigit()) or "1"
            parts.append(f"{'#' * int(level)} {text}")
        else:
            parts.append(text)
    for t_index, table in enumerate(document.tables, start=1):
        parts.append(f"\n[Table {t_index}]")
        for row in table.rows:
            parts.append(" | ".join(cell.text.strip() for cell in row.cells))
    return "\n".join(parts)


def read_document(path: Path) -> str:
    suffix = path.suffix.lower()
    if suffix == ".pdf":
        return _read_pdf(path)
    if suffix == ".docx":
        return _read_docx(path)
    return path.read_text(encoding="utf-8", errors="replace")


def load_documents(folder: Path) -> list[SourceDocument]:
    """Read every supported file under `folder` (recursively), sorted by path for stable output."""
    if not folder.is_dir():
        raise FileNotFoundError(f"Input folder not found: {folder}")
    docs: list[SourceDocument] = []
    for path in sorted(folder.rglob("*")):
        if not path.is_file() or path.suffix.lower() not in SUPPORTED_EXTENSIONS:
            continue
        if any(part.startswith(".") for part in path.relative_to(folder).parts):
            continue
        text = read_document(path).strip()
        if text:
            docs.append(SourceDocument(path=path, name=str(path.relative_to(folder)), text=text))
    return docs


def fingerprint(docs: list[SourceDocument]) -> str:
    """One hash over all inputs; changes whenever any document is added, removed or edited."""
    h = hashlib.sha256()
    for doc in docs:
        h.update(doc.name.encode("utf-8"))
        h.update(doc.sha256.encode("utf-8"))
    return h.hexdigest()


def as_context(docs: list[SourceDocument]) -> str:
    """Render all documents into one tagged block that the model can cite by file name."""
    blocks = [f'<document name="{doc.name}">\n{doc.text}\n</document>' for doc in docs]
    return "<business_requirement_documents>\n" + "\n\n".join(blocks) + "\n</business_requirement_documents>"
